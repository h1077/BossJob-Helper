"""DOCX resume exporter — consumes build_docx_render_plan() contract.

Uses python-docx to produce a clean .docx file from the same semantic
document model that drives HTML preview and PDF export.
"""
from __future__ import annotations

import io
from typing import Any

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from utils.resume_renderer import build_docx_render_plan
from utils.resume_templates import normalize_template_config


def export_resume_docx(
    structured_data: dict[str, Any],
    title: str = "",
    template_config: dict[str, Any] | None = None,
    hidden_modules: list[str] | None = None,
) -> bytes:
    """Return a .docx file as bytes."""
    plan = build_docx_render_plan(structured_data, title, template_config, hidden_modules)
    cfg = plan["config"]
    accent = _parse_color(cfg.get("accentColor", "#2563eb"))
    doc = DocxDocument()

    # Page margins
    margin_mm = cfg.get("pageMargin", 16)
    for section in doc.sections:
        section.top_margin = Inches(margin_mm / 25.4)
        section.bottom_margin = Inches(margin_mm / 25.4)
        section.left_margin = Inches(margin_mm / 25.4)
        section.right_margin = Inches(margin_mm / 25.4)

    body_font_size = Pt(cfg.get("bodyFontSize", 11))

    for section_data in plan["sections"]:
        kind = section_data.get("kind", "")
        section_title = section_data.get("title", "")
        entries = section_data.get("entries", [])

        if kind == "header":
            _add_header(doc, section_data, accent)
        elif kind in ("card_list", "skill_list", "award_list", "text_block"):
            if section_title:
                _add_section_title(doc, section_title, accent, body_font_size)
            for entry in entries:
                if kind == "card_list":
                    _add_card(doc, entry, body_font_size)
                elif kind == "skill_list":
                    _add_skill_group(doc, entry, accent, body_font_size)
                elif kind == "award_list":
                    _add_award_item(doc, entry, body_font_size)
                elif kind == "text_block":
                    _add_text_block(doc, entry, body_font_size)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _parse_color(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip("#")
    try:
        return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))
    except (ValueError, IndexError):
        return RGBColor(0x25, 0x63, 0xEB)


def _add_header(doc, section_data, accent):
    title = section_data.get("title", "")
    items = section_data.get("entries", [])

    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    if items:
        p = doc.add_paragraph()
        for item in items:
            run = p.add_run(f"{item.get('label','')}: {item.get('value','')}  ")
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    # Accent underline
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("_" * 80)
    run.font.color.rgb = accent
    run.font.size = Pt(4)


def _add_section_title(doc, title, accent, body_size):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(body_size.pt * 1.15)
    # Bottom border via thin paragraph rule
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    run2 = p2.add_run("_" * 70)
    run2.font.size = Pt(1)
    run2.font.color.rgb = accent


def _add_card(doc, entry, body_size):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    left = entry.get("left", "")
    center = entry.get("center", "")
    right = entry.get("right", "")

    if left:
        run = p.add_run(left)
        run.bold = True
        run.font.size = Pt(body_size.pt)
    if center:
        run = p.add_run(f"  {center}")
        run.font.size = Pt(body_size.pt - 1)
    if right:
        run = p.add_run(f"  | {right}")
        run.font.size = Pt(body_size.pt - 2)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    sub = entry.get("sub", "")
    if sub:
        for line in str(sub).split("\n"):
            if line.strip():
                p2 = doc.add_paragraph()
                p2.paragraph_format.space_before = Pt(1)
                run = p2.add_run(line.strip())
                run.font.size = Pt(body_size.pt - 1.5)
                run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    tags = entry.get("tags", "")
    if tags:
        p3 = doc.add_paragraph()
        run = p3.add_run(tags)
        run.font.size = Pt(body_size.pt - 2)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)


def _add_skill_group(doc, entry, accent, body_size):
    label = entry.get("label", "")
    tags = entry.get("tags", [])
    desc = entry.get("desc", "")

    if label:
        p = doc.add_paragraph()
        run = p.add_run(label + ": ")
        run.bold = True
        run.font.size = Pt(body_size.pt - 1)

    if tags:
        run = p.add_run("  ".join(str(t) for t in tags))
        run.font.size = Pt(body_size.pt - 1)

    if desc:
        p2 = doc.add_paragraph()
        run = p2.add_run(desc)
        run.font.size = Pt(body_size.pt - 2)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


def _add_award_item(doc, entry, body_size):
    p = doc.add_paragraph()
    label = entry.get("label", "")
    value = entry.get("value", "")
    if label:
        run = p.add_run(label + " ")
        run.bold = True
        run.font.size = Pt(body_size.pt - 1)
    if value:
        run = p.add_run(value + "  ")
        run.font.size = Pt(body_size.pt - 1)


def _add_text_block(doc, entry, body_size):
    value = entry.get("value", "")
    if value:
        p = doc.add_paragraph()
        run = p.add_run(value)
        run.font.size = Pt(body_size.pt - 1)
